"""Module for creating rating curve plots and spreadsheets from HYSTRUC.DAT data."""

import os
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.marker import Marker
from modules.utilities import time_function
from modules.hystruc_extraction import hystruc_extract
import pandas as pd
from typing import Optional, Tuple, Dict

@time_function
def hystruc_rating_table_plots(file_path: str) -> Tuple[Optional[str], Optional[str]]:
    """
    Main function to process HYSTRUC.DAT file and create outputs.
    
    Args:
        file_path: Path to the directory containing HYSTRUC.DAT
        
    Returns:
        Tuple of (pdf_filename, excel_filename) or (None, None) if no rating tables
    """
    hystruc_file = os.path.join(file_path, "HYSTRUC.DAT")
    if not os.path.exists(hystruc_file):
        return None, None
        
    # Extract data from HYSTRUC.DAT
    hystruc_data = hystruc_extract(hystruc_file)
    
    # Create outputs if rating tables exist
    if 'rating_tables' in hystruc_data and not hystruc_data['rating_tables'].empty:
        output_folder = os.path.join(file_path, 'flo2d_plots')
        os.makedirs(output_folder, exist_ok=True)

        pdf_filename = os.path.join(output_folder, 'hystruc_rating_tables.pdf')
        excel_filename = os.path.join(output_folder, 'hystruc_rating_tables.xlsx')

        _plot_rating_tables_to_pdf(hystruc_data['rating_tables'], pdf_filename)
        _create_rating_tables_spreadsheet(hystruc_data['rating_tables'], excel_filename)

        return pdf_filename, excel_filename
        
    return None, None

@time_function
def _plot_rating_tables_to_pdf(rating_tables_df: pd.DataFrame, pdf_filename: str) -> None:
    """
    Plot rating tables (discharge vs depth) to a PDF file with 4 plots per page.

    Args:
        rating_tables_df: DataFrame containing rating table data from hystruc_extract
        pdf_filename: The output file path for the PDF
    """
    with PdfPages(pdf_filename) as pdf:
        # Prepare the plot layout: 2x2 grid on each page
        num_plots_per_page = 4
        fig, axes = plt.subplots(2, 2, figsize=(8.5, 11))
        fig.subplots_adjust(hspace=0.3, wspace=0.3)
        
        # Get unique structures
        structures = rating_tables_df['structure_name'].unique()
        plot_count = 0
        
        for structure in structures:
            # Get data for this structure
            structure_data = rating_tables_df[rating_tables_df['structure_name'] == structure]
            
            # Plot depth vs discharge
            ax = axes[plot_count // 2, plot_count % 2]
            ax.plot(structure_data['discharge'], structure_data['depth'], 
                   color='blue', label='Stage vs Discharge', marker='o')
            ax.set_title(structure)
            ax.set_xlabel('Discharge (cfs)')
            ax.set_ylabel('Stage (ft)')
            ax.grid(True)
            
            # Move to next plot position
            plot_count += 1
            
            # If we've filled the 2x2 grid, save the page and start a new one
            if plot_count == num_plots_per_page:
                pdf.savefig(fig)
                plt.close(fig)
                fig, axes = plt.subplots(2, 2, figsize=(8.5, 11))
                fig.subplots_adjust(hspace=0.3, wspace=0.3)
                plot_count = 0
        
        # Save any remaining plots on the final page
        if plot_count > 0:
            for remaining in range(plot_count, num_plots_per_page):
                axes[remaining // 2, remaining % 2].axis('off')  # Turn off unused subplots
            pdf.savefig(fig)
            plt.close(fig)

@time_function
def _create_rating_tables_spreadsheet(rating_tables_df: pd.DataFrame, excel_filename: str) -> None:
    """
    Create an Excel spreadsheet with rating table data and plots for each structure.

    Args:
        rating_tables_df: DataFrame containing rating table data from hystruc_extract
        excel_filename: The output file path for the Excel spreadsheet
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove the default sheet

    # Get unique structures
    structures = rating_tables_df['structure_name'].unique()

    for structure in structures:
        # Get data for this structure
        structure_data = rating_tables_df[rating_tables_df['structure_name'] == structure]
        
        # Create a new worksheet for each structure
        ws = wb.create_sheet(title=structure)

        # Write headers and data
        ws.append(["Stage (ft)", "Discharge (cfs)"])
        for _, row in structure_data.iterrows():
            ws.append([row['depth'], row['discharge']])

        # Create a scatter plot
        chart = ScatterChart()
        chart.title = f"Rating Table - {structure}"
        chart.x_axis.title = "Discharge (cfs)"
        chart.y_axis.title = "Stage (ft)"

        # Define data ranges for the chart
        x_values = Reference(ws, min_col=2, min_row=2, max_row=len(structure_data) + 1)
        y_values = Reference(ws, min_col=1, min_row=2, max_row=len(structure_data) + 1)

        # Create series
        series = Series(y_values, x_values, title="Stage vs Discharge")
        series.marker = Marker(symbol='circle', size=7)
        series.marker.graphicalProperties.solidFill = "4472C4"  # Blue color
        series.marker.graphicalProperties.line.solidFill = "4472C4"  # Blue outline
        series.graphicalProperties.line.solidFill = "4472C4"  # Blue line
        series.graphicalProperties.line.width = 20000  # Line width
        series.smooth = True  # Smoothed line
        chart.series.append(series)

        # Add the chart to the worksheet
        ws.add_chart(chart, "D2")

    # Save the workbook
    wb.save(excel_filename)

if __name__ == "__main__":
    # Add parent directory to Python path for direct script execution
    import sys
    from pathlib import Path
    sys.path.append(str(Path(__file__).parent.parent))
    
    # Example usage
    folder_path = r"R:\_anichols\Projects\AZ_7_RANCHES\FLO2D\20231212_Added_FPXSEC"
    
    # Process HYSTRUC.DAT and create outputs
    pdf_file, excel_file = hystruc_rating_table_plots(folder_path)
    
    if pdf_file and excel_file:
        print(f"Rating tables plotted and saved to: {pdf_file}")
        print(f"Rating tables spreadsheet created: {excel_file}")
    else:
        print("No rating tables found in HYSTRUC.DAT")