import os
import pandas as pd
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import LineChart, Reference
from modules.swmm_rpt_timeseries_extraction import parse_swmm_nodes_to_dataframe, parse_swmm_links_to_dataframe

def create_timeseries_pdf_plots(df, output_pdf_path, element_type):
    """
    Creates a PDF with time series plots (4 per page) for SWMM elements.
    Matches the style of inlet plots.
    """
    plt.style.use('default')
    
    # Convert time column to datetime if it's not already
    if not pd.api.types.is_datetime64_any_dtype(df['Time']):
        df['Time'] = pd.to_datetime(df['Time'], format='%b-%d-%Y %H:%M:%S')
    
    # Calculate elapsed hours
    start_time = df['Time'].iloc[0]
    df['Elapsed Hours'] = (df['Time'] - start_time).dt.total_seconds() / 3600
    
    with PdfPages(output_pdf_path) as pdf:
        elements = [col for col in df.columns if col not in ['Time', 'Elapsed Hours']]
        num_pages = (len(elements) + 3) // 4

        for page in range(num_pages):
            fig, axs = plt.subplots(2, 2, figsize=(8.5, 11))
            fig.subplots_adjust(hspace=0.4, wspace=0.3)
            axs = axs.flatten()

            for i in range(4):
                idx = page * 4 + i
                if idx >= len(elements):
                    break
                    
                element = elements[idx]
                ax = axs[i]
                
                # Plot data with blue line
                ax.plot(df['Elapsed Hours'], df[element], color='blue', label='Flow')
                
                # Set title
                ax.set_title(f'{element_type} {element}')
                
                # Configure axes labels
                ax.set_xlabel('Time (hours)')
                ax.set_ylabel('Flow (cfs)')
                
                # Configure grid - simpler style matching inlet plots
                ax.grid(True, linestyle='-', color='#E6E6E6')
                
                # Set x-axis limits to match inlet plots
                ax.set_xlim(0, 30)
                
                # Add peak flow information in text box
                peak_flow = df[element].max()
                peak_time_hrs = df['Elapsed Hours'][df[element].idxmax()]
                label = f'Peak Flow: {peak_flow:.2f} cfs\nTime of Peak: {peak_time_hrs:.2f} hrs'
                ax.text(0.05, 0.95, label, 
                       ha='left', 
                       va='top', 
                       transform=ax.transAxes, 
                       fontsize=8,
                       bbox=dict(facecolor='white', alpha=0.6))

            # Remove unused subplots
            for j in range(i + 1, 4):
                fig.delaxes(axs[j])

            pdf.savefig(fig)
            plt.close(fig)

def create_timeseries_excel(df, excel_path, element_type):
    """
    Creates an Excel file with time series data and live plots.
    
    Args:
        df (pd.DataFrame): DataFrame containing time series data
        excel_path (str): Path to save the Excel file
        element_type (str): Type of SWMM element ('Node' or 'Link')
    """
    wb = Workbook()
    wb.remove(wb.active)

    # Create summary sheet
    summary = wb.create_sheet("Summary")
    elements = [col for col in df.columns if col != 'Time']
    
    # Write summary headers
    summary['A1'] = f"{element_type} Name"
    summary['B1'] = "Peak Flow (cfs)"
    summary['C1'] = "Time of Peak"
    
    # Calculate and write summary data
    for idx, element in enumerate(elements, 2):
        peak_flow = df[element].max()
        time_of_peak = df['Time'][df[element].idxmax()]
        summary[f'A{idx}'] = element
        summary[f'B{idx}'] = peak_flow
        summary[f'C{idx}'] = time_of_peak

    # Create data sheet with all time series
    data_sheet = wb.create_sheet("Time Series Data")
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            data_sheet.cell(row=r_idx, column=c_idx, value=value)

    # Create individual sheets with plots for each element
    for element in elements:
        ws = wb.create_sheet(title=element[:31])  # Excel sheet names limited to 31 chars
        
        # Write time and flow data
        ws['A1'] = 'Time'
        ws['B1'] = 'Flow (cfs)'
        for idx, (time, flow) in enumerate(zip(df['Time'], df[element]), 2):
            ws[f'A{idx}'] = time
            ws[f'B{idx}'] = flow

        # Create chart
        chart = LineChart()
        chart.title = f"{element}\nPeak Flow = {df[element].max():.2f} cfs"
        chart.x_axis.title = "Time"
        chart.y_axis.title = "Flow (cfs)"
        
        data = Reference(ws, min_col=2, min_row=1, max_row=len(df) + 1)
        time = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(time)

        ws.add_chart(chart, "D5")

    wb.save(excel_path)

def create_swmm_timeseries_outputs(folder_path, rpt_file):
    """
    Creates PDF and Excel outputs for SWMM nodes and links time series data.
    
    Args:
        folder_path (str): Base folder path
        rpt_file (str): Name of the SWMM report file
    """
    # Create output folder if it doesn't exist
    out_folder = os.path.join(folder_path, 'flo2d_plots')
    os.makedirs(out_folder, exist_ok=True)

    # File paths
    rpt_path = os.path.join(folder_path, rpt_file)
    
    # Process nodes
    df_nodes = parse_swmm_nodes_to_dataframe(rpt_path)
    create_timeseries_pdf_plots(df_nodes, 
                              os.path.join(out_folder, 'swmm_nodes_timeseries.pdf'),
                              'Node')
    create_timeseries_excel(df_nodes,
                           os.path.join(out_folder, 'swmm_nodes_timeseries.xlsx'),
                           'Node')

    # Process links
    df_links = parse_swmm_links_to_dataframe(rpt_path)
    create_timeseries_pdf_plots(df_links,
                               os.path.join(out_folder, 'swmm_links_timeseries.pdf'),
                               'Link')
    create_timeseries_excel(df_links,
                           os.path.join(out_folder, 'swmm_links_timeseries.xlsx'),
                           'Link')

if __name__ == "__main__":
    folder_path = r'R:\_anichols\Projects\_flo2d_postprocessor_tests\Detroit_Basin_Prop100y24h'
    file_path = os.path.join(folder_path, 'swmm.RPT')  # Replace with actual file path
    df_nodes = parse_swmm_nodes_to_dataframe(file_path)
    print(df_nodes.head())
    df_links = parse_swmm_links_to_dataframe(file_path)
    print(df_links.head())
    
    create_swmm_timeseries_outputs(folder_path, 'swmm.RPT')